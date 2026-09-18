"""Exact workbook schemas for the two additional ledger sheets."""
from __future__ import annotations
import csv,io,json,re,uuid,hashlib
from datetime import datetime,date,time,timedelta
from pathlib import Path
from ledger_core import META,now,sha,normalize_time
BIRTH=["生产日期","牛场登记生产时间","牛号","预产期","提前预产期天数","是否佩戴过尾环","是否处于有效监测范围","备注(最后一次佩戴时间）"]
WEARING=["日期","新佩戴牛号","设备编码","拆除时间(掉落）","佩戴时长（天）","设备去向","库存数量","备注"]
MANAGEMENT=["日期","设备来源","设备数量","设备号","故障数量","故障设备编码","故障率(%)","遗失设备","备注"]
COMMON=["记录日期","牧场","工作表","记录类型","核对提示","来源","原始单元格","时间说明","归类目录"]
def text(value):
    if value is None:return ""
    if isinstance(value,datetime):return value.strftime("%Y-%m-%d") if not(value.hour or value.minute or value.second) else value.isoformat(sep=" ")
    if isinstance(value,time):return value.isoformat()
    if isinstance(value,date):return value.isoformat()
    if isinstance(value,float) and value.is_integer():return str(int(value))
    return str(value)
def parse_date(value,year=2026,epoch=None):
    if isinstance(value,datetime):return value.date().isoformat()
    if isinstance(value,date):return value.isoformat()
    if isinstance(value,(int,float)) and 30000<=value<=80000:
        base=epoch or datetime(1899,12,30)
        return (base+timedelta(days=float(value))).date().isoformat()
    value=str(value or "").strip()
    try:
        if len(value)>10:return datetime.fromisoformat(value).date().isoformat()
        return date.fromisoformat(value).isoformat()
    except ValueError:pass
    m=re.fullmatch(r"(?:(\d{4})[年/.-])?(\d{1,2})[月/.-](\d{1,2})[日号]?",value)
    if m:
        try:return date(int(m[1] or year),int(m[2]),int(m[3])).isoformat()
        except ValueError:return ""
    return ""
def parse_removal(value,year=2026,epoch=None):
    if isinstance(value,datetime):return text(value)
    if isinstance(value,(int,float)) and 30000<=value<=80000:
        return text((epoch or datetime(1899,12,30))+timedelta(days=float(value)))
    raw=text(value).strip()
    if len(raw)>10:
        try:return text(datetime.fromisoformat(raw))
        except ValueError:pass
    basic=parse_date(value,year,epoch)
    if basic:return basic
    raw=text(value).strip();m=re.fullmatch(r"(\d{1,2})月(\d{1,2})日(\d{1,2})时(?:(\d{1,2})分?)?",raw)
    if m:
        try:return datetime(year,int(m[1]),int(m[2]),int(m[3]),int(m[4] or 0)).strftime("%Y-%m-%d %H:%M")
        except ValueError:pass
    return raw
def number(value):
    try:return float(value)
    except (TypeError,ValueError):return None
def device_list(value):return re.findall(r"(?<![0-9A-Fa-f])[0-9A-Fa-f]{4,12}(?![0-9A-Fa-f])",str(value or ""))
class SheetSchema:
    def __init__(self,id,title,filename,business,date_field,cow_field):
        self.id=id;self.title=title;self.filename=filename;self.business=business;self.fields=business+COMMON+META;self.date_field=date_field;self.cow_field=cow_field
        self.defaults={"sheet_id":id,"server_file":str(Path(r"F:\牛舍_现场记录")/filename)}
    def empty(self,settings=None,section=None):
        row=dict.fromkeys(self.fields,"");row.update({"记录ID":str(uuid.uuid4()),"已删除":"0","记录日期":datetime.now().strftime("%Y-%m-%d"),"牧场":(settings or {}).get("farm","扬大_高邮牧场"),"工作表":section or ("产犊登记" if self.id=="calving" else "佩戴台账"),"记录类型":"数据"})
        row[self.date_field]=row["记录日期"];return row
    def normalize(self,row):
        result={f:text(row.get(f,"")) for f in self.fields}
        result["记录日期"]=parse_date(result.get(self.date_field)) or result.get("记录日期","")
        result["已删除"]=result["已删除"] or "0";result["归类目录"]=self.directory(result)
        return result
    def directory(self,row):
        return "/".join([row.get("牧场",""),self.title,row.get("工作表",""),row.get("记录日期") or "日期待补充",row.get(self.cow_field) or row.get("记录类型","")])
    def fingerprint(self,row):
        values={f:text(row.get(f,"")) for f in self.business+["记录日期","牧场","工作表","记录类型","已删除"]}
        values["已删除"]=values["已删除"] or "0"
        return sha(json.dumps(values,ensure_ascii=False,sort_keys=True,separators=(",",":")).encode())
    def identity(self,row):
        when=row.get("记录日期","");cow=row.get(self.cow_field,"")
        if self.id=="calving":parts=[row.get("牧场",""),when,cow]
        elif row.get("工作表")=="佩戴台账":parts=[row.get("牧场",""),when,cow,row.get("设备编码","").upper()]
        else:return ""
        if not when or not cow or not parts[-1] or parts[-1] in ("?","？"):return ""
        return sha(json.dumps([self.id]+parts,ensure_ascii=False).encode())
    def issues(self,row,required=True):
        errors=[]
        if row.get("记录类型") in ("日期分组","汇总"):return errors
        when=row.get("记录日期","")
        if not parse_date(when):errors.append("记录日期待核对")
        if self.id=="calving":
            if not row.get("牛号"):errors.append("牛号未填写")
            clock=row.get("牛场登记生产时间","")
            if not re.fullmatch(r"\d{2}:\d{2}(:\d{2})?",clock):errors.append("牛场登记生产时间待核对")
            due=parse_date(row.get("预产期"));early=number(row.get("提前预产期天数"))
            if due and parse_date(when) and early is not None:
                actual=(date.fromisoformat(due)-date.fromisoformat(when)).days
                if early!=actual:errors.append(f"提前预产期天数原填{row['提前预产期天数']}，按日期计算应为{actual}，请核实")
            for field in ("是否佩戴过尾环","是否处于有效监测范围"):
                if row.get(field,"").strip() not in ("","是","否","/"):errors.append(field+"取值待核对")
        elif row.get("工作表")=="佩戴台账":
            if row.get("记录类型")=="佩戴":
                if not row.get("新佩戴牛号"):errors.append("新佩戴牛号未填写")
                device=row.get("设备编码","")
                if not re.fullmatch(r"[0-9A-Fa-f]{4}|[0-9A-Fa-f]{12}",device):errors.append("设备编码长度或字符待核对")
                removal=row.get("拆除时间(掉落）","")
                if removal in ("?","？"):errors.append("拆除时间(掉落）待核对")
                if removal and removal not in ("?","？","/"):
                    end=parse_date(removal)
                    if not end:errors.append("拆除时间(掉落）格式待核对")
                    elif when and end<when:errors.append("拆除时间(掉落）早于佩戴日期")
                    elif len(removal)==10 and number(row.get("佩戴时长（天）")) is not None and when:
                        days=(date.fromisoformat(end)-date.fromisoformat(when)).days
                        if number(row["佩戴时长（天）"])!=days:errors.append(f"佩戴时长（天）原填{row['佩戴时长（天）']}，两日期相差{days}天，需核对计时口径")
        elif row.get("工作表")=="设备管理":
            for count,devices in [("设备数量","设备号"),("故障数量","故障设备编码")]:
                n=number(row.get(count));codes=device_list(row.get(devices))
                if n is not None and codes and n!=len(codes):errors.append(f"{count}原填{row[count]}，{devices}列有{len(codes)}个编码，请核实")
                if any(len(c) not in (4,12) for c in codes):errors.append(devices+"存在非4位或12位编码")
            total=number(row.get("设备数量"));bad=number(row.get("故障数量"));rate=number(row.get("故障率(%)"))
            if total is not None and bad is not None:
                if bad>total:errors.append("故障数量大于设备数量")
                if total and rate is not None and abs(rate-bad/total*100)>0.0001:errors.append("故障率(%)与数量计算不符")
        return errors
    def csv_bytes(self,rows):
        stream=io.StringIO(newline="");writer=csv.DictWriter(stream,fieldnames=self.fields,lineterminator="\r\n");writer.writeheader()
        for row in sorted(rows,key=lambda r:(r.get("记录日期",""),r.get("工作表",""),r.get("来源",""),r.get("记录ID",""))):writer.writerow({f:row.get(f,"") for f in self.fields})
        return stream.getvalue().encode("utf-8-sig")
    def read_csv(self,data):
        reader=csv.DictReader(io.StringIO(data.decode("utf-8-sig"),newline=""),strict=True)
        if reader.fieldnames!=self.fields:raise ValueError("CSV表头不属于"+self.title)
        result=list(reader)
        if any(None in r or any(v is None for v in r.values()) for r in result):raise ValueError("CSV行列数不匹配")
        return result
SCHEMAS={
    "calving":SheetSchema("calving","产犊登记","扬大产犊登记汇总.csv",BIRTH,"生产日期","牛号"),
    "equipment":SheetSchema("equipment","设备台账","扬大测试设备台账.csv",list(dict.fromkeys(WEARING+MANAGEMENT)),"日期","新佩戴牛号")
}
def get_schema(id="samples"):
    if id=="samples":return None
    if id not in SCHEMAS:raise ValueError("未知台账Sheet")
    return SCHEMAS[id]
def detect_workbook(path):
    import openpyxl
    wb=openpyxl.load_workbook(path,read_only=True,data_only=False)
    try:
        for ws in wb:
            h=[text(ws.cell(2,c).value).strip() for c in range(1,10)]
            if h[:8]==BIRTH:return "calving"
            if h[:8]==WEARING:return "equipment"
            if [text(ws.cell(3,c).value).strip() for c in (1,2,3)]==["序号","牛号","预产期"]:return "samples"
    finally:wb.close()
    raise ValueError("表头与三张台账模板均不匹配，已停止导入")
def import_workbook(path,settings=None):
    import openpyxl
    id=detect_workbook(path)
    if id=="samples":
        from ledger_import import import_excel
        return id,import_excel(path,settings)
    schema=SCHEMAS[id];wb=openpyxl.load_workbook(path,data_only=False);cached=openpyxl.load_workbook(path,data_only=True)
    rows=[];year_candidates=[]
    for ws in wb:
        for r in range(3,ws.max_row+1):
            for c in (1,4) if id=="calving" else (1,):
                value=ws.cell(r,c).value
                if isinstance(value,datetime):year_candidates.append(value.year)
                elif isinstance(value,(int,float)) and 30000<=value<=80000:year_candidates.append(int(parse_date(value,epoch=wb.epoch)[:4]))
    year=max(set(year_candidates),key=year_candidates.count) if year_candidates else None
    if not year:raise ValueError("原表缺少可确认年份，无法补齐日期；请在表内填写完整年份")
    for ws in wb:
        headers=BIRTH if id=="calving" else WEARING if ws.title=="佩戴台账" else MANAGEMENT if ws.title=="设备管理" else None
        if not headers:
            if any(c.value is not None for row in ws for c in row):raise ValueError("存在未识别的工作表："+ws.title)
            continue
        if [text(ws.cell(2,c).value).strip() for c in range(1,len(headers)+1)]!=headers:raise ValueError(ws.title+"表头不完全匹配")
        if any(ws.cell(r,c).value is not None for r in range(1,ws.max_row+1) for c in range(len(headers)+1,ws.max_column+1)):raise ValueError(ws.title+"表头范围外有数据，请核对后再导入")
        merged={}
        for m in ws.merged_cells.ranges:
            if m.min_row<3:continue
            for r in range(m.min_row,m.max_row+1):
                for c in range(m.min_col,m.max_col+1):merged[r,c]=(m.min_row,m.min_col)
        date_context="";date_origin=""
        for r in range(3,ws.max_row+1):
            physical=[ws.cell(r,c) for c in range(1,len(headers)+1)]
            if not any(c.value is not None for c in physical):continue
            row=schema.empty(settings,section="产犊登记" if id=="calving" else ws.title);notes=[];raw={}
            for c,key in enumerate(headers,1):
                origin=merged.get((r,c),(r,c));cell=ws.cell(*origin);original=ws.cell(r,c);value=cell.value
                record={"field":key,"value":text(original.value),"type":type(original.value).__name__,"format":original.number_format,"resolved_from":cell.coordinate,"font_color":str(cell.font.color) if cell.font.color else "","fill_color":str(cell.fill.fgColor) if cell.fill.patternType else "","font_rgb":cell.font.color.rgb if cell.font.color and cell.font.color.type=="rgb" else "","comment":cell.comment.text if cell.comment else ""}
                if cell.data_type=="f":
                    cached_value=cached[ws.title][cell.coordinate].value
                    record.update(formula=cell.value,cached=text(cached_value))
                    if cached_value is None:raise ValueError(f"{ws.title}!{cell.coordinate}公式没有缓存结果，请在Excel计算保存后导入")
                    value=cached_value
                raw[original.coordinate]=record
                row[key]=text(value)
                if key in ("生产日期","预产期","日期"):
                    normalized=parse_date(value,year,wb.epoch)
                    if normalized:row[key]=normalized
                if key=="拆除时间(掉落）":row[key]=parse_removal(value,year,wb.epoch)
                if origin!=(r,c):notes.append(original.coordinate+"沿用合并单元格"+cell.coordinate)
            current=parse_date(row[schema.date_field],year)
            if current:date_context=current;date_origin=f"{ws.title}!A{r}"
            if id=="calving":
                if not row["生产日期"]:
                    row["生产日期"]=date_context;notes.append("生产日期沿用"+date_origin)
                row["记录类型"]="产犊"
            elif ws.title=="佩戴台账":
                if not row["日期"]:row["日期"]=date_context;notes.append("日期沿用"+date_origin)
                elif not current:notes.append("日期栏原文保留："+row["日期"]+"；日期归属沿用"+date_origin)
                row["记录类型"]="佩戴" if row["设备编码"] else "库存/说明" if any(row.get(f) for f in ("库存数量","备注","设备去向")) else "日期分组"
            else:
                row["记录类型"]="汇总" if row["日期"]=="合计" or row["设备来源"]=="目前库存" else "设备批次" if row["设备数量"] else "管理说明"
            row["记录日期"]=date_context if row["记录类型"]!="汇总" else ""
            row["来源"]=Path(path).name+" / "+ws.title+" / 第"+str(r)+"行"
            row["原始单元格"]=json.dumps(raw,ensure_ascii=False,separators=(",",":"))
            row["时间说明"]="；".join(notes)
            row["归类目录"]=schema.directory(row);row["核对提示"]="；".join(schema.issues(row))
            row["记录ID"]=str(uuid.uuid5(uuid.NAMESPACE_URL,id+":"+schema.fingerprint(row)))
            rows.append(row)
    wb.close();cached.close()
    return id,rows
