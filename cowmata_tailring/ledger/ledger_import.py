"""Import the user workbook without executing formulas or document instructions."""
import calendar, json, re, uuid
from datetime import datetime
from pathlib import Path
from .ledger_core import *
def legacy_time(value,year,anchor=None,date_only=False):
    if value is None: return ""
    if isinstance(value,datetime): return value.strftime("%Y-%m-%d" if date_only else "%Y-%m-%d %H:%M")
    text=str(value).strip()
    text=re.sub(r"^(大约|约)", "", text)
    if text in ("","/"): return text
    try: return normalize_time(text,date_only)
    except ValueError: pass
    if date_only: return text
    m=re.fullmatch(r"(?:(\d{1,2})月)?(\d{1,2})日(\d{1,2})时(?:(\d{1,2})分?)?",text)
    if not m: return text
    month,day,hour,minute=m.groups();minute=int(minute or 0)
    try:
        if month:
            candidate=datetime(year,int(month),int(day),int(hour),minute)
            if anchor and anchor not in ("","/"):
                try:
                    ref=datetime.fromisoformat(anchor)
                    candidates=[candidate.replace(year=y) for y in [ref.year-1,ref.year,ref.year+1]]
                    candidate=min(candidates,key=lambda d:abs((d-ref).total_seconds()))
                except ValueError: pass
            return candidate.strftime("%Y-%m-%d %H:%M")
        if anchor and anchor not in ("","/"):
            ref=datetime.fromisoformat(anchor)
            candidates=[]
            for offset in (-1,0,1):
                n=ref.year*12+ref.month-1+offset;y,mo=divmod(n,12)
                try: candidates.append(datetime(y,mo+1,int(day),int(hour),minute))
                except ValueError: pass
            candidates.sort(key=lambda d:abs((d-ref).total_seconds()))
            if len(candidates)==1 or abs((candidates[0]-ref).total_seconds())<abs((candidates[1]-ref).total_seconds()):
                return candidates[0].strftime("%Y-%m-%d %H:%M")
    except (ValueError,OverflowError): pass
    return text
def _mismatch(sheet,row,expected,label):
    from .ledger_sheets import canonical_header
    from openpyxl.utils import get_column_letter
    bad=[f"{get_column_letter(c)}{row}应为「{v}」实际「{str(sheet.cell(row,c).value or '').strip() or '空'}」" for c,v in expected.items() if canonical_header(sheet.cell(row,c).value)!=canonical_header(v)]
    return f"{sheet.title}{label}与样本试验台账不一致，未导入：" + "；".join(bad) if bad else ""
def import_excel(path,settings=None):
    import openpyxl
    from .ledger_sheets import canonical_header
    book=openpyxl.load_workbook(path,data_only=False)
    result=[]
    try:
        for sheet in book:
            header=None
            for n in range(1,min(sheet.max_row,20)+1):
                cells=[canonical_header(sheet.cell(n,c).value) for c in range(1,16)]
                if cells[:4]==[canonical_header(v) for v in ("序号","牛号","预产期","佩戴记录")]: header=n;break
            if not header: continue
            main_expected={1:'序号',2:'牛号',3:'预产期',4:'佩戴记录',7:'产犊时间段',9:'监测目的',10:'样本有效性预判',13:'尾环佩戴人',14:'事件标注人',15:'样本评价'}
            problem=_mismatch(sheet,header,main_expected,'主表头')
            if problem:raise ValueError(problem)
            expected={4:"设备号",5:"佩戴开始",6:"佩戴结束",7:"开始",8:"结束",10:"九轴",11:"脉搏",12:"温度"}
            problem=_mismatch(sheet,header+1,expected,'子表头')
            if problem:raise ValueError(problem)
            title=" ".join(str(sheet.cell(n,1).value or "") for n in range(1,header))
            years=re.findall(r"(20\d{2})年",title)
            if not years: raise ValueError("台账标题没有年份，无法确定简写时间；请补全标题年份")
            year=int(years[0])
            merged={}
            for area in sheet.merged_cells.ranges:
                if area.min_row<=header+1: continue
                for r in range(area.min_row,area.max_row+1):
                    for c in range(area.min_col,min(area.max_col,15)+1):
                        merged[(r,c)]=sheet.cell(area.min_row,area.min_col).value
            for n in range(header+2,sheet.max_row+1):
                direct=[sheet.cell(n,c).value for c in range(1,16)]
                if not any(v is not None for v in direct): continue
                values=[merged.get((n,c),sheet.cell(n,c).value) for c in range(1,16)]
                row=empty_record(settings);raw={f:("" if v is None else str(v)) for f,v in zip(BUSINESS,values)}
                row.update(raw)
                row["预产期"]=legacy_time(values[2],year,date_only=True)
                row["佩戴开始"]=legacy_time(values[4],year)
                row["佩戴结束"]=legacy_time(values[5],year, row["佩戴开始"])
                row["产犊开始"]=legacy_time(values[6],year,row["佩戴开始"])
                try:
                    datetime.fromisoformat(row["产犊开始"]);end_anchor=row["产犊开始"]
                except ValueError: end_anchor=row["佩戴开始"]
                row["产犊结束"]=legacy_time(values[7],year,end_anchor)
                row["时间说明"]="；".join(f+"为估计时间（"+raw[f]+"）" for f in ["佩戴开始","佩戴结束","产犊开始","产犊结束"] if re.match(r"^(大约|约)",raw[f]))
                try: row["记录日期"]=datetime.fromisoformat(row["佩戴开始"]).date().isoformat()
                except ValueError: row["记录日期"]=""
                row["数据分类"]=category_for(row["监测目的"])
                row["原始单元格"]=json.dumps(raw,ensure_ascii=False,separators=(",",":"))
                row["来源"]=f"{Path(path).name} / {sheet.title} / 第{n}行"
                identity=json.dumps([sheet.title,raw],ensure_ascii=False,sort_keys=True)
                row["记录ID"]=str(uuid.uuid5(uuid.NAMESPACE_URL,"cowmata-ledger-import:"+identity))
                row["核对提示"]="；".join(issues(row))
                row["归类目录"]=directory(row)
                result.append(row)
    finally: book.close()
    if not result: raise ValueError("未找到样本试验台账表头或数据")
    return result
