"""Surgical XLSX cell updates: preserve workbook parts, styles, merges and formulas."""
import io,re,zipfile,copy,ast,operator
import xml.etree.ElementTree as ET
from datetime import datetime,date,time
from pathlib import PurePosixPath
from .ledger_core import BUSINESS
from .ledger_sheets import layout_for
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';N={'m':NS};REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
ET.register_namespace('',NS)
def tag(name):return '{'+NS+'}'+name
def colnum(value):
    total=0
    for c in value:total=total*26+ord(c)-64
    return total

def patch_workbook(original,store,edits,fresh):
    import openpyxl
    from openpyxl.utils import get_column_letter,range_boundaries
    from openpyxl.utils.datetime import to_excel
    book=openpyxl.load_workbook(io.BytesIO(original),data_only=False)
    patches={};appended=[]
    try:
        layouts={}
        def layout(sheet):
            if sheet not in layouts:layouts[sheet]=layout_for(book[sheet],store.schema.id)
            return layouts[sheet]
        def columns_for(sheet):
            """Field -> column from the file's own header row, so reordered or renamed layouts write back in place."""
            if not store.schema:return {key:c for c,key in enumerate(BUSINESS,1)}
            found=layout(sheet)
            if found is None:raise ValueError(sheet+' 表头已无法识别，未回写原文件；请核对表头后重新导入')
            return dict(found.columns)
        def sheet_for(section):
            if section in book.sheetnames:return section
            return next((name for name in book.sheetnames if (layout(name) and layout(name).canonical_title==section)),section)
        # Merged values can only be replaced when every linked row agrees; no implicit format changes.
        allrows={r['记录ID']:r for r in store.all(True)}
        def put(sheet,coordinate,value):
            key=(sheet,coordinate)
            if key in patches and patches[key]!=value:raise ValueError(sheet+'!'+coordinate+' 的合并单元格存在不同修改，请先在 Excel 核对')
            patches[key]=value
        from .ledger_files import extra_values,field_value,same_cell
        for (sheet,n),baseline,row in edits:
            if sheet not in book.sheetnames:raise ValueError('原工作表不存在：'+sheet)
            ws=book[sheet]
            if row.get('已删除')=='1':raise ValueError('原 Excel 不含删除标记，请在 Excel 中核对删除；本次未上传')
            import json
            columns=columns_for(sheet);columns.update(json.loads(row.get('原始单元格') or '{}').get('__extra_columns__',{}))
            for key,c in columns.items():
                if same_cell(field_value(row,key),field_value(baseline,key)):continue
                coord=f'{get_column_letter(c)}{n}'
                for area in ws.merged_cells.ranges:
                    if coord in area:
                        for link in store.db.execute('SELECT slot,record_id FROM source_rows'):
                            import json
                            other_sheet,other_n=json.loads(link['slot'])
                            other=allrows.get(link['record_id'])
                            if other_sheet==sheet and area.min_row<=other_n<=area.max_row and other and not same_cell(field_value(other,key),field_value(row,key)):
                                raise ValueError(f'{sheet}!{area} 为合并单元格，相关行必须填写相同值，或先在 Excel 拆分后重新导入')
                        coord=ws.cell(area.min_row,area.min_col).coordinate;break
                put(sheet,coord,field_value(row,key))
        next_rows={ws.title:max((c.row for row in ws for c in row if c.value is not None),default=2)+1 for ws in book}
        for row in fresh:
            sheet=sheet_for(row.get('工作表')) if store.schema and store.schema.id=='equipment' else book.sheetnames[0]
            if sheet not in book.sheetnames:raise ValueError('新增记录的工作表不存在')
            n=next_rows[sheet];next_rows[sheet]+=1
            for key,c in columns_for(sheet).items():put(sheet,f'{get_column_letter(c)}{n}',row.get(key,''))
            appended.append(((sheet,n),row))
        # Retain cell types where they are meaningful; text identifiers remain text.
        typed={}
        for (sheet,coord),value in patches.items():
            cell=book[sheet][coord];old=cell.value;value=str(value)
            if cell.data_type=='f':raise ValueError(f'{sheet}!{coord} 是原表公式，请修改输入列，程序不会覆盖公式')
            converted=value
            if isinstance(old,(datetime,date,time)):
                try:
                    converted=to_excel(datetime.fromisoformat(value),book.epoch)
                except ValueError:
                    try:converted=to_excel(time.fromisoformat(value),book.epoch)
                    except ValueError:pass
            elif isinstance(old,(float,int)) and not isinstance(old,bool) and re.fullmatch(r'-?\d+(?:\.\d+)?',value):converted=float(value)
            typed[sheet,coord]=converted
        with zipfile.ZipFile(io.BytesIO(original)) as archive:
            workbook=ET.fromstring(archive.read('xl/workbook.xml'));rels=ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
            targets={r.get('Id'):r.get('Target') for r in rels if r.get('TargetMode')!='External'}
            parts={s.get('name'):(targets[s.get('{'+REL+'}id')].lstrip('/') if targets[s.get('{'+REL+'}id')].startswith('/') else 'xl/'+targets[s.get('{'+REL+'}id')]) for s in workbook.find('m:sheets',N)}
            replacements={}
            for sheet in {s for s,c in patches}:
                part=parts[sheet]
                if '..' in PurePosixPath(part).parts:raise ValueError('工作表路径无效')
                xml=ET.fromstring(archive.read(part));data=xml.find('m:sheetData',N);nodes={c.get('r'):c for r in data for c in r};rownodes={int(r.get('r')):r for r in data}
                changed=set()
                for (s,coord),value in typed.items():
                    if s!=sheet:continue
                    n=int(re.search(r'\d+',coord)[0]);c=nodes.get(coord)
                    if c is None:
                        rownode=rownodes.get(n)
                        if rownode is None:rownode=ET.SubElement(data,tag('row'),{'r':str(n)});rownodes[n]=rownode
                        c=ET.SubElement(rownode,tag('c'),{'r':coord});nodes[coord]=c
                        # Copy the nearest preceding row style only, never its contents/formula.
                        prefix=re.match('[A-Z]+',coord)[0]
                        for prev in range(n-1,1,-1):
                            oldnode=nodes.get(prefix+str(prev))
                            if oldnode is not None and oldnode.get('s') is not None:c.set('s',oldnode.get('s'));break
                    for child in list(c):
                        if child.tag in (tag('v'),tag('is'),tag('f')):c.remove(child)
                    c.attrib.pop('t',None)
                    if isinstance(value,(float,int)):
                        ET.SubElement(c,tag('v')).text=format(value,'.15g')
                    elif value!='':
                        c.set('t','inlineStr');t=ET.SubElement(ET.SubElement(c,tag('is')),tag('t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=value
                    changed.add(coord)
                recalculate(xml,changed)
                for r in data:r[:]=sorted(r,key=lambda c:colnum(re.match('[A-Z]+',c.get('r'))[0]))
                data[:]=sorted(data,key=lambda r:int(r.get('r')))
                dim=xml.find('m:dimension',N)
                if dim is not None:
                    maxcol=max(colnum(re.match('[A-Z]+',c)[0]) for c in nodes);maxrow=max(int(re.search(r'\d+',c)[0]) for c in nodes)
                    dim.set('ref',f'A1:{get_column_letter(maxcol)}{maxrow}')
                replacements[part]=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
            output=io.BytesIO()
            with zipfile.ZipFile(output,'w') as dest:
                for item in archive.infolist():dest.writestr(item,replacements.get(item.filename,archive.read(item.filename)))
            return output.getvalue(),appended
    finally:book.close()

def recalculate(xml,changed):
    """Update affected simple formulas from the supplied template, never evaluate arbitrary code."""
    from openpyxl.utils.cell import range_boundaries,get_column_letter
    cells={c.get('r'):c for c in xml.findall('.//m:sheetData/m:row/m:c',N)};active=set();done={};ops={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv}
    def number(coord):
        if coord in done:return done[coord]
        if coord in active:raise ValueError('公式循环引用：'+coord)
        c=cells.get(coord)
        if c is None:return 0
        formula=c.find('m:f',N)
        if formula is not None:
            active.add(coord)
            expr=formula.text or '';dependencies=set(re.findall(r'\$?[A-Z]+\$?\d+',expr));dependencies={d.replace('$','') for d in dependencies}
            def sumrange(m):
                a,b,c2,d=range_boundaries(m[1].replace('$',''));coords=[f'{get_column_letter(col)}{row}' for row in range(b,d+1) for col in range(a,c2+1)];dependencies.update(coords)
                return str(sum(number(x) for x in coords))
            expr=re.sub(r'SUM\((\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+)\)',sumrange,expr,flags=re.I)
            values={ref:number(ref) for ref in dependencies}
            affected=bool(dependencies&changed)
            if not affected:
                active.remove(coord);v=c.find('m:v',N);return float(v.text or 0) if v is not None else 0
            expr=re.sub(r'\$?([A-Z]+)\$?(\d+)',lambda m:str(values[m[1]+m[2]]),expr)
            def calc(node):
                if isinstance(node,ast.Expression):return calc(node.body)
                if isinstance(node,ast.Constant) and type(node.value) in (int,float):return node.value
                if isinstance(node,ast.BinOp) and type(node.op) in ops:return ops[type(node.op)](calc(node.left),calc(node.right))
                if isinstance(node,ast.UnaryOp) and isinstance(node.op,ast.USub):return -calc(node.operand)
                raise ValueError('原表包含无法安全重算的公式，请在 Excel 保存后重新导入：'+coord)
            try:value=calc(ast.parse(expr,mode='eval'))
            except (SyntaxError,ZeroDivisionError) as e:raise ValueError('公式计算待核对：'+coord) from e
            v=c.find('m:v',N)
            if v is None:v=ET.SubElement(c,tag('v'))
            v.text=format(value,'.15g');c.attrib.pop('t',None);changed.add(coord);done[coord]=value;active.remove(coord);return value
        if c.get('t') in ('inlineStr','s','str'):return 0
        v=c.find('m:v',N)
        try:return float(v.text or 0) if v is not None else 0
        except ValueError:return 0
    for coord,c in cells.items():
        if c.find('m:f',N) is not None:number(coord)
