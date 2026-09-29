"""Local source-file links; scans whole files but only changes differing records."""
import csv,io,json,re,uuid
from pathlib import Path
from .ledger_core import Store,FIELDS,BUSINESS,sha,empty_record,issues,read_csv
from .ledger_sheets import SCHEMAS,import_workbook

class SourceFileError(ValueError):
    def __init__(self,path,message):
        self.path=str(path);super().__init__(message+'：'+self.path)

def slot(row):
    match=re.fullmatch(r"(.+) / (.+) / 第(\d+)行",row.get("来源",""))
    return (match[2],int(match[3])) if match else (row.get("工作表","CSV"),0)
def init(store):
    if store.db.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name IN ('source_files','source_rows','source_conflicts')").fetchone()[0]==3:return
    store.db.executescript("""CREATE TABLE IF NOT EXISTS source_files(path TEXT PRIMARY KEY,document TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS source_rows(path TEXT NOT NULL,slot TEXT NOT NULL,record_id TEXT NOT NULL,body TEXT NOT NULL,PRIMARY KEY(path,slot));
CREATE TABLE IF NOT EXISTS source_conflicts(path TEXT NOT NULL,record_id TEXT NOT NULL,incoming TEXT NOT NULL,PRIMARY KEY(path,record_id));""")
def business(store):return store.schema.business if store.schema else BUSINESS
def same_cell(a,b):return str(a or "").strip()==str(b or "").strip()
def same_business(store,a,b):return all(same_cell(field_value(a,k),field_value(b,k)) for k in business(store)+["已删除"]+list(set(extra_values(a))|set(extra_values(b))))
def read_document(path,settings=None):
    path=Path(path).resolve();data=path.read_bytes();document={"path":str(path),"sha256":sha(data),"format":path.suffix.lower().lstrip('.'),"notices":[]}
    if path.suffix.lower()==".xlsx":
        id,rows=import_workbook(path,settings,document["notices"])
        enrich_display(path,id,rows,document["notices"])
    elif path.suffix.lower()==".csv":
        encoding='utf-8-sig'
        try:content=data.decode(encoding)
        except UnicodeDecodeError:encoding='gb18030';content=data.decode(encoding)
        reader=csv.DictReader(io.StringIO(content,newline=''),strict=True);headers=reader.fieldnames
        if not headers or len(set(headers))!=len(headers):raise ValueError('CSV 表头为空或含重复列名，未导入')
        from .ledger_sheets import canonical_header
        # Header identity ignores width, bracket style and stray spaces; column order stays the contract for CSV.
        known={canonical_header(f):f for f in FIELDS+['配种开始','配种结束']+[f for s in SCHEMAS.values() for f in s.fields]}
        header_map={h:known.get(canonical_header(h),h) for h in headers}
        if len(set(header_map.values()))!=len(headers):raise ValueError('CSV 表头规范化后出现重复列名，未导入')
        original_headers=headers;headers=[header_map[h] for h in headers]
        if headers!=original_headers:document['notices'].append('CSV 表头已按标准列名识别：'+'；'.join(f'{a}→{b}' for a,b in header_map.items() if a!=b))
        base_headers=[h for h in headers if h not in ('配种开始','配种结束')]
        id='samples' if base_headers in (FIELDS,BUSINESS) else next((id for id,s in SCHEMAS.items() if headers in (s.fields,s.business)),None)
        if id is None:raise ValueError('CSV 表头与三个台账均不匹配；请保留原表列名和顺序')
        schema=SCHEMAS.get(id);rows=[]
        for n,values in enumerate(reader,2):
            if None in values or any(v is None for v in values.values()):raise ValueError(f'CSV 第{n}行列数不符')
            values={header_map[k]:v for k,v in values.items()}
            row=(schema.empty(settings) if schema else empty_record(settings));row.update(values)
            if '记录ID' not in headers:
                row['来源']=f'{path.name} / CSV / 第{n}行'
                row['原始单元格']=json.dumps(values,ensure_ascii=False)
                row['记录ID']=str(uuid.uuid5(uuid.NAMESPACE_URL,'cowmata-business-csv:'+json.dumps(values,ensure_ascii=False,sort_keys=True)))
                if not schema:
                    from .ledger_core import category_for
                    row['数据分类']=category_for(row['监测目的']);row['记录日期']=row['佩戴开始'][:10]
                else:row=schema.normalize(row)
            if id=='samples' and any(k in headers for k in ('配种开始','配种结束')):
                raw=json.loads(row.get('原始单元格') or '{}');raw['__extra__']={k:values[k] for k in headers if k in ('配种开始','配种结束')};raw['__base_time_note__']=row.get('时间说明','');row['原始单元格']=json.dumps(raw,ensure_ascii=False)
                row['时间说明']='；'.join([row.get('时间说明','')]+[k+'='+v for k,v in raw['__extra__'].items()]).strip('；')
            rows.append(row)
        document.update(headers=original_headers,header_map=header_map,encoding=encoding,newline='\r\n' if '\r\n' in content else '\n',canonical='记录ID' in headers)
    else:raise ValueError('请选择 XLSX 或 CSV 文件')
    if path.read_bytes()!=data:raise ValueError('读取过程中原文件发生变化，请保存文件后重新导入')
    for row in rows:
        if id=='samples':
            from .ledger_core import classify_record,directory
            code,reason=classify_record(row);row['数据分类']=code;row['归类目录']=directory(row)
        uuid.UUID(row['记录ID']);row['核对提示']='；'.join((SCHEMAS[id].issues if id in SCHEMAS else issues)(row))
    document.update(sheet_id=id,rows=rows)
    return document

def import_document(store,document):
    store.reconcile_conflicts()
    init(store);path=str(Path(document['path']).resolve());incoming=[dict(r) for r in document['rows']]
    current={r['记录ID']:r for r in store.all(True)}
    bindings={tuple(json.loads(r['slot'])):dict(r) for r in store.db.execute('SELECT * FROM source_rows WHERE path=?',(path,))}
    prior={r['record_id']:json.loads(r['body']) for r in bindings.values()}
    matched={};used=set();ambiguous=set();ambiguous_candidates={}
    from collections import defaultdict
    fingerprints=defaultdict(list);events=defaultdict(list);serials=defaultdict(list);positions=defaultdict(list)
    for key,r in current.items():
        fingerprints[store.fingerprint(r)].append(key)
        if store.identity(r):events[store.identity(r)].append(key)
        worksheet,_=slot(r)
        if key in prior or r.get("来源", "").split(" / ")[0]==Path(path).name:serials[(worksheet,r.get("序号", ""))].append(key)
        if r.get("来源", "").split(" / ")[0]==Path(path).name:positions[slot(r)].append(key)
    def choose(i,candidates):
        candidates=set(candidates)-used
        if len(candidates)==1:
            key=next(iter(candidates));matched[i]=key;used.add(key)
        elif len(candidates)>1:ambiguous.add(i);ambiguous_candidates.setdefault(i,set()).update(candidates)
    # Match all unchanged/moved records before using positions, so inserted rows do not displace identities.
    for i,row in enumerate(incoming):
        if row['记录ID'] in current:choose(i,[row['记录ID']])
    for i,row in enumerate(incoming):
        if i not in matched:
            choose(i,fingerprints[store.fingerprint(row)])
    for i,row in enumerate(incoming):
        if i in matched:continue
        event=store.identity(row)
        if event:choose(i,events[event])
    from .ledger_matching import continuation
    for i,row in enumerate(incoming):
        if i not in matched:choose(i,[r['记录ID'] for r in continuation(row,current.values(),store.schema)])
    # A sample serial is only usable within this same source worksheet, and when unique.
    for i,row in enumerate(incoming):
        if i in matched:continue
        worksheet,_=slot(row)
        if not store.schema and row.get('序号'):
            candidates=serials[(worksheet,row['序号'])]
            choose(i,candidates)
    for i,row in enumerate(incoming):
        if i in matched:continue
        position=slot(row);binding=bindings.get(position)
        if i in ambiguous and not binding:continue
        candidates=[binding['record_id']] if binding else positions[position]
        choose(i,candidates)
    from .ledger_merge import merge_record,DeletedRecordConflict
    # A renamed daily edition retains the previous per-cell baseline.
    from .ledger_matching import family
    old_links=list(store.db.execute('SELECT rowid,* FROM source_rows WHERE path!=? ORDER BY rowid DESC',(path,)))
    for link in old_links:
        if family(link['path'])==family(path) and link['record_id'] in matched.values():
            prior.setdefault(link['record_id'],json.loads(link['body']))
    allowed=set();rows=[];local_only=set();file_rows={}
    for i,row in enumerate(incoming):
        if i in ambiguous and i not in matched:continue
        if i in matched:
            key=matched[i];row['记录ID']=key;file_rows[key]=dict(row)
            if key in prior and same_business(store,row,prior[key]) and not same_business(store,current[key],row):
                local_only.add(key)
            else:
                try:
                    row,_,_=merge_record(prior.get(key),row,current[key],store.schema)
                    allowed.add(key)
                except DeletedRecordConflict:pass
        else:file_rows[row['记录ID']]=dict(row)
        rows.append(row)
    store.import_rows([r for r in rows if r['记录ID'] not in local_only],review_updates=False,allow_dirty_ids=allowed)
    summary=dict(store.last_import_summary);summary['skipped']+=len(local_only)
    ambiguous_rows=[incoming[i].get('来源','') for i in sorted(ambiguous-set(matched))]
    ids=[]
    with store.db:
        store.db.execute('DELETE FROM source_conflicts WHERE path=?',(path,))
        for merged_row in rows:
            row=file_rows[merged_row['记录ID']]
            local=store.get(row['记录ID'])
            if not local:continue
            if row['记录ID'] in local_only:
                store.db.execute('DELETE FROM source_rows WHERE path=? AND record_id=?',(path,row['记录ID']))
                store.db.execute('INSERT OR REPLACE INTO source_rows VALUES(?,?,?,?)',(path,json.dumps(slot(row),ensure_ascii=False),row['记录ID'],json.dumps(row,ensure_ascii=False)))
                continue
            if not same_business(store,local,merged_row):
                store.db.execute('INSERT OR REPLACE INTO source_conflicts VALUES(?,?,?)',(path,row['记录ID'],json.dumps(row,ensure_ascii=False)))
                continue
            ids.append(row['记录ID'])
            if local.get('原始单元格')!=row.get('原始单元格') or local.get('来源')!=row.get('来源'):
                stored={k:local.get(k,'') for k in store.fields};stored.update({'原始单元格':row.get('原始单元格',''),'来源':row.get('来源','')})
                store.db.execute('UPDATE records SET body=? WHERE id=?',(json.dumps(stored,ensure_ascii=False),row['记录ID']))
            position=slot(row)
            store.db.execute('DELETE FROM source_rows WHERE path=? AND record_id=?',(path,row['记录ID']))
            store.db.execute('INSERT OR REPLACE INTO source_rows VALUES(?,?,?,?)',(path,json.dumps(position,ensure_ascii=False),row['记录ID'],json.dumps(row,ensure_ascii=False)))
        metadata={k:v for k,v in document.items() if k!='rows'}
        metadata['ambiguous_rows']=ambiguous_rows
        metadata['ambiguous_details']=[{'source':incoming[i].get('来源',''),'incoming':incoming[i],'candidates':sorted(ambiguous_candidates.get(i,set()))} for i in sorted(ambiguous-set(matched))]
        summary['conflicts']=store.db.execute('SELECT COUNT(*) FROM source_conflicts WHERE path=?',(path,)).fetchone()[0]+len(ambiguous_rows)
        metadata['unresolved_conflicts']=summary['conflicts']
        metadata['managed_csv']=Path(path).resolve() in (store.csv_path.resolve(),Path(store.settings()['server_file']).resolve())
        store.db.execute('INSERT OR REPLACE INTO source_files VALUES(?,?)',(path,json.dumps(metadata,ensure_ascii=False)))
    # Accepted rows now write back to this edition, never to every historical copy.
    with store.db:
        for old in store.db.execute('SELECT * FROM source_files WHERE path!=?',(path,)).fetchall():
            if family(old['path'])!=family(path):continue
            links=store.db.execute('SELECT * FROM source_rows WHERE path=?',(old['path'],)).fetchall()
            retired=[dict(link) for link in links if link['record_id'] in set(ids)|local_only]
            meta=json.loads(old['document'])
            if not retired:
                # Older releases could leave only an unresolved counter after their row map was lost.
                # Once a same-family daily file imports cleanly, archive that orphan marker instead of
                # blocking every subsequent import with a legacy review item. No records are deleted.
                if not links and meta.get('unresolved_conflicts') and not ambiguous_rows and not summary.get('conflicts'):
                    meta['superseded_by']=path;meta['unresolved_conflicts']=0
                    meta['ambiguous_rows']=[];meta['ambiguous_details']=[]
                    meta['legacy_archived_without_rows']=True
                    store.db.execute('UPDATE source_files SET document=? WHERE path=?',(json.dumps(meta,ensure_ascii=False),old['path']))
                continue
            meta.setdefault('superseded_links',[]).extend(retired)
            for link in retired:
                store.db.execute('DELETE FROM source_rows WHERE path=? AND slot=?',(old['path'],link['slot']))
                store.db.execute('DELETE FROM source_conflicts WHERE path=? AND record_id=?',(old['path'],link['record_id']))
            if len(retired)==len(links) and not meta.get('ambiguous_rows'):
                meta['superseded_by']=path;meta['unresolved_conflicts']=0
            store.db.execute('UPDATE source_files SET document=? WHERE path=?',(json.dumps(meta,ensure_ascii=False),old['path']))
    store.configure({'source_file':path})
    store.last_import_summary=summary
    return ids


def writeback(store,record_ids=None):
    """Write local edits to linked sources first. Never replace a source with the server union."""
    from .ledger_core import atomic_write
    init(store);plans=[];linked=set();pending={p['record']['记录ID']:p['record'] for p in store.pending() if record_ids is None or p['record']['记录ID'] in record_ids}
    for entry in store.db.execute('SELECT * FROM source_files'):
        path=Path(entry['path']);meta=json.loads(entry['document'])
        if meta.get('managed_csv') or meta.get('superseded_by'):continue
        bindings=list(store.db.execute('SELECT * FROM source_rows WHERE path=?',(str(path),)))
        edits=[]
        for link in bindings:
            linked.add(link['record_id']);row=pending.get(link['record_id']);baseline=json.loads(link['body'])
            if row and not same_business(store,row,baseline):edits.append((tuple(json.loads(link['slot'])),baseline,row))
        fresh=[]
        if str(path)==store.settings().get('source_file'):
            all_linked={r[0] for r in store.db.execute('SELECT record_id FROM source_rows')}
            fresh=[r for key,r in pending.items() if key not in all_linked and r.get('已删除')!='1' and not r.get('来源')]
        if not edits and not fresh:continue
        if meta.get('unresolved_conflicts'):raise SourceFileError(path,'原文件还有待核对内容，请在下方核对面板处理；文件已移动可点击重新选择原文件')
        if not path.is_file():raise SourceFileError(path,'找不到关联原文件，请点击下方重新选择原文件；本地修改保留')
        original=path.read_bytes()
        if sha(original)!=meta['sha256']:raise SourceFileError(path,'原文件刚刚发生变化，请点击下方后台重新核对')
        if meta['format']=='xlsx':
            from .ledger_workbook import patch_workbook
            content,appended=patch_workbook(original,store,edits,fresh)
        else:
            content,appended=patch_csv(original,store,meta,edits,fresh)
        plans.append((path,meta,original,content,edits,appended))
    # Validate every file before the first replacement.
    for path,meta,original,content,edits,appended in plans:
        if path.read_bytes()!=original:raise ValueError('原文件正在被其他程序修改，请稍后再同步：'+str(path))
    for path,meta,original,content,edits,appended in plans:
        backup=store.root/'source-backups'/(sha(original)+'-'+path.name)
        if not backup.exists():atomic_write(backup,original)
        if path.read_bytes()!=original:raise ValueError('原文件已变化，未覆盖：'+str(path))
        atomic_write(path,content)
        meta['sha256']=sha(content)
        with store.db:
            store.db.execute('UPDATE source_files SET document=? WHERE path=?',(json.dumps(meta,ensure_ascii=False),str(path)))
            for position,baseline,row in edits:
                store.db.execute('UPDATE source_rows SET body=? WHERE path=? AND slot=?',(json.dumps(row,ensure_ascii=False),str(path),json.dumps(position,ensure_ascii=False)))
            for position,row in appended:
                store.db.execute('INSERT OR REPLACE INTO source_rows VALUES(?,?,?,?)',(str(path),json.dumps(position,ensure_ascii=False),row['记录ID'],json.dumps(row,ensure_ascii=False)))
    for path,*_ in plans:
        import_document(store,read_document(path,store.settings()))
    return [str(plan[0]) for plan in plans]

def patch_csv(original,store,meta,edits,fresh):
    headers=meta['headers'];names=meta.get('header_map') or {h:h for h in headers}
    reader=csv.DictReader(io.StringIO(original.decode(meta.get('encoding','utf-8-sig')),newline=''),strict=True);rows=list(reader)
    for position,baseline,row in edits:
        if meta.get('canonical'):
            matches=[i for i,r in enumerate(rows) if r['记录ID']==baseline['记录ID']]
        else:matches=[position[1]-2] if 0<=position[1]-2<len(rows) else []
        if len(matches)!=1:raise ValueError('CSV 原始行无法唯一定位，请重新导入')
        target=rows[matches[0]]
        if row.get('已删除')=='1':
            column=next((h for h in headers if names.get(h,h)=='已删除'),None)
            if column:target[column]='1'
            else:raise ValueError('业务 CSV 不含删除标记，请在原文件中删除行后重新导入；本次未上传')
        else:
            for column in headers:
                key=names.get(column,column)
                if key in business(store)+list(extra_values(row)) and not same_cell(field_value(row,key),field_value(baseline,key)):target[column]=field_value(row,key)
    appended=[]
    for row in fresh:
        rows.append({column:field_value(row,names.get(column,column)) for column in headers});appended.append((('CSV',len(rows)+1),row))
    out=io.StringIO(newline='');w=csv.DictWriter(out,fieldnames=headers,lineterminator=meta.get('newline','\r\n'));w.writeheader();w.writerows(rows)
    return out.getvalue().encode(meta.get('encoding','utf-8-sig')),appended


def cell_display(cell,cached=None):
    from datetime import datetime,date,time
    value=cached.value if cell.data_type=='f' and cached is not None else cell.value
    if value is None:return ''
    if isinstance(value,str):return value
    fmt=cell.number_format.split(';')[0]
    if isinstance(value,(datetime,date,time)):
        if '年' in fmt:return f'{value.year}年{value.month:02d}月{value.day:02d}日'
        if isinstance(value,time) or ('h' in fmt.lower() and 'y' not in fmt.lower() and 'd' not in fmt.lower()):
            return (f'{value.hour:02d}' if 'hh' in fmt.lower() else str(value.hour))+f':{value.minute:02d}'+(f':{value.second:02d}' if 's' in fmt.lower() else '')
        if fmt.lower()=='mm-dd-yy':return f'{value.month:02d}-{value.day:02d}-{value.year%100:02d}'
        if 'y' in fmt.lower():return f'{value.year}/{value.month:02d}/{value.day:02d}'
        return f'{value.month}/{value.day}'
    if isinstance(value,(int,float)):
        if re.search(r'0\.0',fmt):return f'{value:.1f}'
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)

def enrich_display(path,id,rows,notices=None):
    import openpyxl
    from .ledger_sheets import layout_for,canonical_header,used_width
    from openpyxl.utils import get_column_letter
    notices=[] if notices is None else notices;helper_columns={}
    wb=openpyxl.load_workbook(path,data_only=False);cached=openpyxl.load_workbook(path,data_only=True)
    layouts={}
    try:
        for row in rows:
            sheet,n=slot(row);ws=wb[sheet]
            extras={}
            if id=='samples':
                header_row=next(n for n in range(1,min(ws.max_row,20)+1) if canonical_header(ws.cell(n,1).value)=='序号' and canonical_header(ws.cell(n,4).value)=='佩戴记录')
                breeding={canonical_header(k):k for k in ('配种开始','配种结束')}
                for c in range(16,used_width(ws)+1):
                    header=str(ws.cell(header_row,c).value or ws.cell(header_row+1,c).value or '').strip()
                    if canonical_header(header) in breeding and breeding[canonical_header(header)] not in extras:extras[breeding[canonical_header(header)]]=c
                    elif ws.cell(n,c).value is not None:
                        # Helper columns are kept verbatim instead of blocking the whole workbook.
                        raw_extra=json.loads(row.get('原始单元格') or '{}');raw_extra.setdefault('__unmapped__',{})[ws.cell(n,c).coordinate]={'header':header,'value':str(ws.cell(n,c).value)}
                        row['原始单元格']=json.dumps(raw_extra,ensure_ascii=False,separators=(',',':'))
                        helper_columns.setdefault(sheet,{})[c]=header
            if id=='samples':columns={field:c for c,field in enumerate(BUSINESS,1)}
            else:
                if sheet not in layouts:layouts[sheet]=layout_for(ws,id)
                columns=layouts[sheet].columns if layouts[sheet] else {}
            display={field:{'value':cell_display(ws.cell(n,c),cached[sheet].cell(n,c)),'normalized':row.get(field,'')} for field,c in columns.items()}
            raw=json.loads(row.get('原始单元格') or '{}');raw['__display__']=display
            if extras:
                from .ledger_import import legacy_time
                raw['__extra_columns__']=extras
                raw['__extra__']={key:legacy_time(ws.cell(n,c).value,int(row.get('记录日期','')[:4] or 2026),row.get('佩戴开始')) for key,c in extras.items()}
                for key,c in extras.items():display[key]={'value':cell_display(ws.cell(n,c),cached[sheet].cell(n,c)),'normalized':raw['__extra__'][key]}
                raw['__base_time_note__']=row.get('时间说明','')
                row['时间说明']='；'.join([row.get('时间说明','')]+[key+'='+raw['__extra__'][key] for key in extras]).strip('；')
            row['原始单元格']=json.dumps(raw,ensure_ascii=False,separators=(',',':'))
    finally:wb.close();cached.close()
    for sheet,columns in helper_columns.items():
        notices.append(f"{sheet}非标准列 "+"、".join(f"{get_column_letter(c)}列「{h or '无表头'}」" for c,h in sorted(columns.items()))+" 的原值已保留在原始单元格，不参与字段比对")

def display_value(row,key):
    try:meta=json.loads(row.get('原始单元格') or '{}').get('__display__',{}).get(key)
    except (ValueError,AttributeError):meta=None
    value=field_value(row,key)
    # Merged Excel cells (one cow, several wearings) keep only the first physical cell;
    # show the inherited value instead of a blank that looks like missing data.
    original=meta['value'] if meta and value==meta.get('normalized') and (meta.get('value') or not value) else value
    # Keep intentional merged/blank cells, but never show an Excel date serial.
    if original=='':return ''
    date_fields={'日期','生产日期','预产期','记录日期'}
    time_fields={'佩戴开始','佩戴结束','产犊开始','产犊结束','拆除时间(掉落）','配种开始','配种结束'}
    if key in date_fields|time_fields:
        from datetime import datetime
        try:
            parsed=datetime.fromisoformat(value)
            return parsed.strftime('%Y-%m-%d' if key in date_fields else '%Y-%m-%d %H:%M')
        except (ValueError,TypeError):pass
    return original


def file_conflicts(store,path):
    init(store)
    return [json.loads(r[0]) for r in store.db.execute('SELECT incoming FROM source_conflicts WHERE path=?',(str(Path(path).resolve()),))]

def resolve_file_conflict(store,path,record_id,use_file):
    init(store);path=str(Path(path).resolve())
    found=store.db.execute('SELECT incoming FROM source_conflicts WHERE path=? AND record_id=?',(path,record_id)).fetchone()
    if not found:raise ValueError('文件冲突已处理，请刷新')
    incoming=json.loads(found[0]);local=store.get(record_id)
    if local.get('_conflict'):raise ValueError('请先处理该记录的服务器版本冲突')
    with store.db:
        if use_file:store._save(incoming,False)
        store.db.execute('INSERT OR REPLACE INTO source_rows VALUES(?,?,?,?)',(path,json.dumps(slot(incoming),ensure_ascii=False),record_id,json.dumps(incoming,ensure_ascii=False)))
        store.db.execute('DELETE FROM source_conflicts WHERE path=? AND record_id=?',(path,record_id))
        entry=store.db.execute('SELECT document FROM source_files WHERE path=?',(path,)).fetchone();meta=json.loads(entry[0]);meta['unresolved_conflicts']=len(file_conflicts(store,path))+len(meta.get('ambiguous_rows',[]))
        store.db.execute('UPDATE source_files SET document=? WHERE path=?',(json.dumps(meta,ensure_ascii=False),path))
    store.export()


def extra_values(row):
    try:return json.loads(row.get('原始单元格') or '{}').get('__extra__',{})
    except (ValueError,AttributeError):return {}
def field_value(row,key):return extra_values(row).get(key,row.get(key,''))
def edit_extra(row,key,value):
    raw=json.loads(row.get('原始单元格') or '{}');raw.setdefault('__extra__',{})[key]=value
    row['原始单元格']=json.dumps(raw,ensure_ascii=False,separators=(',',':'))
    row['时间说明']='；'.join([raw.get('__base_time_note__','')]+[k+'='+v for k,v in raw['__extra__'].items()]).strip('；')
    return row


def source_conflict_count(store):
    init(store)
    known=dict(store.db.execute('SELECT c.path,COUNT(*) FROM source_conflicts c JOIN records r ON r.id=c.record_id GROUP BY c.path'))
    count=store.db.execute('SELECT COUNT(*) FROM source_conflicts c JOIN records r ON r.id=c.record_id WHERE r.conflict IS NULL').fetchone()[0]
    for entry in store.db.execute('SELECT path,document FROM source_files'):
        meta=json.loads(entry['document']);ambiguous=len(meta.get('ambiguous_rows',[]))
        count+=ambiguous+int(meta.get('unresolved_conflicts',0)>known.get(entry['path'],0)+ambiguous)
    return count

def conflict_sources(store):
    init(store)
    return [r['path'] for r in store.db.execute('SELECT * FROM source_files') if json.loads(r['document']).get('unresolved_conflicts') or file_conflicts(store,r['path'])]

def reconcile_sources(store,checkpoint=None):
    """Rescan changed sources before writeback; unchanged files cannot conflict with a local-only edit."""
    init(store);active=store.settings().get('source_file');results=[]
    try:
        for entry in list(store.db.execute('SELECT * FROM source_files')):
            meta=json.loads(entry['document']);path=Path(entry['path'])
            if checkpoint:checkpoint('检查原文件：'+path.name)
            if meta.get('managed_csv') or meta.get('superseded_by') or not path.is_file():continue
            if sha(path.read_bytes())!=meta['sha256'] or meta.get('unresolved_conflicts') or file_conflicts(store,path):
                document=read_document(path,store.settings())
                retired={tuple(json.loads(link['slot'])) for link in meta.get('superseded_links',[])}
                if retired:
                    document['rows']=[r for r in document['rows'] if slot(r) not in retired]
                    document['superseded_links']=meta['superseded_links']
                expected=store.schema.id if store.schema else 'samples'
                if document['sheet_id']!=expected:raise ValueError('原文件已变成其他类型的台账，请重新选择文件：'+str(path))
                if checkpoint:checkpoint('合并单元格：'+path.name)
                import_document(store,document);results.append(str(path))
    finally:
        if active:store.configure({'source_file':active})
    return results


def relocate_source(store,old_path,new_path):
    """Relink a moved source after validating its type and record identity."""
    init(store);old_path=str(Path(old_path).resolve());new_path=str(Path(new_path).resolve())
    entry=store.db.execute('SELECT document FROM source_files WHERE path=?',(old_path,)).fetchone()
    if not entry:raise ValueError('原文件关联已变化，请重新打开核对面板')
    if new_path!=old_path and store.db.execute('SELECT 1 FROM source_files WHERE path=?',(new_path,)).fetchone():
        raise ValueError('所选文件已经关联到此台账，请选择原文件的新位置')
    document=read_document(new_path,store.settings())
    if document['sheet_id']!=(store.schema.id if store.schema else 'samples'):raise ValueError('所选文件的台账类型不匹配，未改变原关联')
    links=list(store.db.execute('SELECT record_id,body FROM source_rows WHERE path=?',(old_path,)))
    ids={r['record_id'] for r in links};bodies=[json.loads(r['body']) for r in links]
    identities={store.identity(r) for r in bodies if store.identity(r)};fingerprints={store.fingerprint(r) for r in bodies}
    if links and not any(r['记录ID'] in ids or store.identity(r) in identities or store.fingerprint(r) in fingerprints for r in document['rows']):
        raise ValueError('所选文件与原台账记录不匹配，未改变原关联；请确认选择的是同一份原文件')
    active=store.settings().get('source_file');meta=json.loads(entry[0]);meta['path']=new_path
    with store.db:
        store.db.execute('UPDATE source_files SET path=?,document=? WHERE path=?',(new_path,json.dumps(meta,ensure_ascii=False),old_path))
        store.db.execute('UPDATE source_rows SET path=? WHERE path=?',(new_path,old_path))
        store.db.execute('UPDATE source_conflicts SET path=? WHERE path=?',(new_path,old_path))
    try:import_document(store,document)
    finally:
        if active:store.configure({'source_file':new_path if active==old_path else active})
    return new_path
